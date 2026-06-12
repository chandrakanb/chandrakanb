import openpyxl

def process_excel(input_file, output_file):
    """
    Processes an Excel file to restructure data based on summary and labels.
    Prints the current summary being processed.

    Args:
        input_file (str): Path to the input Excel file.
        output_file (str): Path to the output Excel file.
    """

    try:
        # Load the workbook
        workbook = openpyxl.load_workbook(input_file)
        sheet = workbook.active

        # Get header row
        header = [cell.value for cell in sheet[1]]
        summary_column_index = header.index("Summary")  # Case-sensitive!
        label_columns = [i for i, col in enumerate(header) if col.startswith("Labels")]

        # Prepare data for the new Excel file
        output_data = []
        for row_num in range(2, sheet.max_row + 1):  # Start from row 2 (assuming row 1 is header)
            summary_value = sheet.cell(row=row_num, column=summary_column_index + 1).value
            if summary_value is None:
                continue  # Skip rows where summary is empty

            print(f"Processing Summary: {summary_value}")  # Print the current summary

            for label_column_index in label_columns:
                label_value = sheet.cell(row=row_num, column=label_column_index + 1).value
                if label_value is not None and str(label_value).strip() != "":  # Check for empty or whitespace-only values
                    output_data.append([summary_value, label_value])

        # Create a new workbook and sheet
        output_workbook = openpyxl.Workbook()
        output_sheet = output_workbook.active
        output_sheet.title = "Processed Data"

        # Write header to the output sheet
        output_sheet.append(["Summary", "Label"])  # Or whatever column names you prefer

        # Write data to the output sheet
        for row in output_data:
            output_sheet.append(row)

        # Save the output workbook
        output_workbook.save(output_file)
        print(f"Successfully processed '{input_file}' and saved to '{output_file}'")

    except FileNotFoundError:
        print(f"Error: Input file '{input_file}' not found.")
    except ValueError as e:
        print(f"Error: {e}")
    except Exception as e:
        print(f"An unexpected error occurred: {e}")

# Example usage:
input_excel_file = "input.xlsx"  # Replace with your input file name
output_excel_file = "output.xlsx"  # Replace with your desired output file name
process_excel(input_excel_file, output_excel_file)

# This is AI generated code, please refer KPIT AI Policy before using this in your projects