import os
import re
import bs4
from bs4 import BeautifulSoup
from datetime import datetime, timedelta
import pandas as pd
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, Border, Side

def read_html_file(execution_report_details):
    html_file = os.path.join(execution_report_details['execution_report_path'], execution_report_details['html_file'])
    df_summary = pd.read_html(html_file)[0]
    df_summary = df_summary.filter(df_summary.columns[:2])

    with open(html_file, 'r', encoding='utf-8') as infile:
        soup = BeautifulSoup(infile, 'html.parser')

    reportframe_add = soup.find('div', attrs={'id': 'reportframe'})
    content_src = None
    for content in reportframe_add.contents:
        if isinstance(content, bs4.Tag):
            content_src = content.get('src')
            break

    content_src = content_src.replace("../", "")
    curr_path = os.path.normpath(os.path.join(execution_report_details['execution_report_path'], os.pardir, 'XMLReport', 'Segments', 'XMLReportSegment_0.xml'))

    tree = ET.parse(curr_path)
    mainlist = []
    for tcs_node in tree.iter("TestCases"):
        for tc_node in tcs_node.iter("TestCase"):
            maindict = {
                "Test ID": tc_node.find("testCaseID").text,
                "Test Objective": tc_node.find("testCaseObjective").text,
                "Start Time": tc_node.find("startTime").text,
                "Video Logger Link": tc_node.find("VideoLogLink").text,
                "Status": tc_node.find("result").text,
                "Total Time": tc_node.find("totalExecutionTime").text,
                "End Time": tc_node.find("endTime").text
            }
            mainlist.append(maindict)

    df_detail_report = pd.DataFrame(mainlist)
    return df_summary, df_detail_report


def iterate_tuple(dir_tuple):
    for item_tuple in dir_tuple:
        if isinstance(item_tuple, tuple):
            iterate_tuple(item_tuple)
        elif isinstance(item_tuple, list):
            for html_file in item_tuple:
                if html_file == "HTMLReport.html":
                    execution_report_path = dir_tuple[0]
                    execution_report_name = execution_report_path.replace(os.getcwd(), "")
                    sub_folder_list = execution_report_name.split('\\')
                    test_pack = None
                    for each_folder in sub_folder_list:
                        if "ITERATION" in each_folder.upper():
                            x = re.search("Iteration_[0-9]+", each_folder)
                            if x:
                                test_pack = each_folder.replace(f"_{x.group()}", "")
                    Bench_name = sub_folder_list[1] if len(sub_folder_list) > 1 else "Unknown"
                    execution_report_name = sub_folder_list[2] if len(sub_folder_list) > 2 else "Unknown"
                    return {
                        "execution_report_name": execution_report_name,
                        "execution_report_path": execution_report_path,
                        "Bench_name": Bench_name,
                        "Test_Pack": test_pack,
                        "html_file": html_file
                    }
    return None


def add_additional_columns(df, execution_report_details):
    for col_ in ['Bench_name', 'execution_report_name', 'Test_Pack']:
        df[col_] = execution_report_details[col_]
    return df


def apply_filter(sheet, keyword, column):
    for cell in sheet[column]:
        if isinstance(keyword, str):
            sheet.auto_filter.add_filter_column(cell.col_idx - 1, [f"*{keyword}*"])
        else:
            sheet.auto_filter.add_filter_column(cell.col_idx - 1, [keyword])


def update_column_B(sheet, keyword, value):
    for cell in sheet.iter_rows(min_row=2, min_col=2, max_row=sheet.max_row, max_col=2):
        if cell[0].offset(column=1).value and isinstance(keyword, str) and keyword.upper() in str(cell[0].offset(column=1).value).upper():
            cell[0].value = value


def pivot_table(sheet):
    data = sheet.values
    cols = next(data)
    data = list(data)
    script_type_idx = cols.index("Script Type")
    status_idx = cols.index("Status")

    counts = {}
    for row in data:
        if row[script_type_idx] not in counts:
            counts[row[script_type_idx]] = {}
        if row[status_idx] not in counts[row[script_type_idx]]:
            counts[row[script_type_idx]][row[status_idx]] = 0
        counts[row[script_type_idx]][row[status_idx]] += 1

    df = pd.DataFrame(counts).fillna(0)
    df.loc['Total'] = df.sum()
    df['Total'] = df.sum(axis=1)
    df.reset_index(inplace=True)

    return df, dataframe_to_rows(df, index=False, header=True), cols


def create_pivot_table(excel_file_path):
    wb = openpyxl.load_workbook(excel_file_path)
    sheet = wb.active
    pivot_sheet = wb.create_sheet(title="temp")
    pivot, pivot_rows, _ = pivot_table(sheet)

    for r_idx, row in enumerate(pivot_rows, 1):
        for c_idx, value in enumerate(row, 1):
            pivot_sheet.cell(row=r_idx, column=c_idx, value=value)

    transposed_pivot_sheet = wb.create_sheet(title="Pivot_Table")
    transposed_pivot = pivot.transpose()

    for r_idx, row in enumerate(dataframe_to_rows(transposed_pivot, index=True, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = transposed_pivot_sheet.cell(row=r_idx, column=c_idx, value=value)

            # Apply border to all cells
            cell.border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            # Apply bold to 1st row and 1st column
            if r_idx == 3 or c_idx == 1:
                cell.font = Font(bold=True)

    transposed_pivot_sheet.delete_rows(1, 2)
    transposed_pivot_sheet['A1'] = "Script Type"
    wb.remove(wb["temp"])
    wb.save(excel_file_path)


if __name__ == "__main__":
    df_summary_main = pd.DataFrame([])
    df_detail_report_main = pd.DataFrame([])

    for dir_tuple in os.walk(os.getcwd()):
        execution_report_details = iterate_tuple(dir_tuple)
        if execution_report_details:
            try:
                df_summary, df_detail_report = read_html_file(execution_report_details)
                df_summary = df_summary.T
                df_summary.columns = df_summary.loc['Summary'].tolist()
                df_summary = add_additional_columns(df_summary, execution_report_details)
                df_detail_report = add_additional_columns(df_detail_report, execution_report_details)

                df_summary_main = pd.concat([df_summary_main, df_summary], axis=0) if not df_summary_main.empty else df_summary
                df_detail_report_main = pd.concat([df_detail_report_main, df_detail_report], axis=0) if not df_detail_report_main.empty else df_detail_report
            except Exception as e:
                print(f"Error processing: {execution_report_details}\n{e}")

    df_summary_main = df_summary_main.loc[df_summary_main['Start Time'] != 'Start Time']
    df_detail_report_main['Start Time_1'] = pd.to_datetime(df_detail_report_main['Start Time'], format="%d %b %Y %H:%M:%S")
    df_detail_report_main['Time_1'] = pd.to_datetime(df_detail_report_main['Total Time'], format="%H:%M:%S")

    End_Time = []
    for x, y in zip(df_detail_report_main['Start Time_1'], df_detail_report_main['Time_1']):
        __day__ = 0 if y.year == 1900 else y.day
        End_Time.append(x + timedelta(days=__day__, hours=y.hour, minutes=y.minute, seconds=y.second))

    df_detail_report_main['End Time'] = End_Time
    df_detail_report_main = df_detail_report_main.filter(["Test_Pack", "Test ID", "Status", "Total Time", "Bench_name"])
    df_detail_report_main_table = df_detail_report_main.copy()
    df_detail_report_main_table['Remark Failure Reason'] = ""

    excel_file_path = "Execution_Report.xlsx"
    with pd.ExcelWriter(excel_file_path, engine='xlsxwriter') as writer:
        df_detail_report_main_table.to_excel(writer, sheet_name='Detailed Report', index=False)

    workbook = openpyxl.load_workbook(excel_file_path)
    sheet = workbook.active
    sheet.insert_cols(2)
    cell = sheet.cell(row=1, column=2)
    cell.value = "Script Type"
    
    cell.font = Font(bold=True)
    cell.border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    for keyword, label in [("KPIT", "Test Script"),
                           ("pre", "Precondition"),
                           ("post", "Postcondition"),
                           ("AutoFlashing", "Precondition")]:
        apply_filter(sheet, keyword, "C")
        update_column_B(sheet, keyword, label)
        sheet.auto_filter.ref = None

    workbook.save(excel_file_path)
    create_pivot_table(excel_file_path)
